"""The export writers (spec 2026-09-23-volumes §10): labels, CSV, XLSX, QML, plan image, PDF."""

import csv
import io
import xml.etree.ElementTree as ET
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook
from PIL import Image as PILImage
from surfaces import CX, CY, WKT, circle, cone, fixture_spec, plane, write_surface

from app.surfaces.grid import open_surface
from app.volumes import report_pdf, tables
from app.volumes.items import ExportItem, labels
from app.volumes.plan_image import render_plan_image
from app.volumes.qml import write_qml

REF = {
    "id": "s1",
    "name": "April survey",
    "kind": "cloud_dsm",
    "method": "median",
    "cell_size_m": 0.1,
    "captured_on": "2026-04-02",
    "cloud_file": "april.las",
    "cloud_sha256": "ab" * 32,
}
RESULTS = {
    "fill_m3": 523.61,
    "cut_m3": 0.04,
    "net_m3": 523.57,
    "unshifted": None,
    "area_m2": 452.1,
    "polygon_area_m2": 452.39,
    "measured_area_m2": 450.0,
    "masked_area_m2": 12.5,
    "excluded_area_m2": 0.0,
    "nodata_area_m2": 2.1,
    "cell_size_m": 0.1,
    "areal_scale_factor": 0.99962,
    "shift_applied_m": 0.0,
    "alignment": None,
    "base_fit": {
        "kind": "toe_plane",
        "samples": 754,
        "rejected": 3,
        "usable_edge_fraction": 0.98,
        "rms_m": 0.012,
        "plane": [0.02, -0.013, 50.9],
    },
    "uncertainty": {
        "total_m3": 6.2,
        "base_m3": 5.4,
        "alignment_m3": None,
        "cell_size_m3": 0.3,
        "nodata_m3": 2.4,
        "patch_m3": 1.1,
        "complete": True,
    },
    "warnings": [{"code": "base_fit_poor", "severity": "warn", "message": "the toe <departs> & more"}],
    "footprints_used": 1,
    "patch_regions": 1,
    "diff_scale_m": 4.2,
    "top_surface": REF,
    "base_surface": None,
    "inputs": {},
    "inputs_fingerprint": "f" * 64,
    "engine_version": 1,
    "computed_at": "2026-09-24T10:00:00+00:00",
    "duration_s": 1.2,
}


def _item(name="Pile & 1", **over) -> ExportItem:
    return ExportItem(
        id="m1",
        name=name,
        status="ready",
        polygon=circle(CX, CY, 12.0, 32),
        base={"kind": "toe_plane"},
        masks={"detection_run_ids": ["r1"], "buffer_m": 1.0, "exclusion_polygons": []},
        alignment={"stable_polygon": None, "apply_shift": False},
        results={**RESULTS, **over},
        epsg=32639,
        crs_wkt=WKT,
    )


def test_labels_follow_the_base_kind():
    assert labels("toe_plane").fill == "Stockpile volume (above base)" and labels("flat").headline == "fill"
    survey = labels("surface", {"kind": "cloud_dsm", "captured_on": "2026-03-01"})
    assert survey.fill == "Added since 2026-03-01 (fill)" and survey.headline == "net"
    design = labels("surface", {"kind": "design", "captured_on": None})
    assert (
        design.fill == "Above design (to cut)"
        and design.cut == "Below design (to fill)"
        and design.headline == "both"
    )


def test_csv_has_the_exact_columns(tmp_path):
    tables.write_csv(tmp_path / "v.csv", [_item()])
    rows = list(csv.DictReader((tmp_path / "v.csv").open(encoding="utf-8")))
    assert list(rows[0]) == tables.CSV_COLUMNS
    assert rows[0]["fill_m3"] == "523.61" and rows[0]["warnings"] == "base_fit_poor"
    assert float(rows[0]["centroid_lat"]) > 29 and rows[0]["epsg"] == "32639"


def test_xlsx_has_three_sheets_with_typed_numbers(tmp_path):
    tables.write_xlsx(tmp_path / "v.xlsx", [_item()])
    wb = load_workbook(tmp_path / "v.xlsx")
    assert wb.sheetnames == ["Volumes", "Method", "Warnings"]
    ws = wb["Volumes"]
    col = tables.CSV_COLUMNS.index("fill_m3") + 1
    cell = ws.cell(row=2, column=col)
    assert cell.value == pytest.approx(523.61) and cell.number_format == "#,##0.0"
    assert ws.freeze_panes == "A2"
    assert wb["Warnings"].cell(row=2, column=2).value == "base_fit_poor"


def test_qml_parses_with_the_symmetric_ramp(tmp_path):
    write_qml(tmp_path / "c.qml", 4.2)
    root = ET.parse(tmp_path / "c.qml").getroot()
    values = [float(i.get("value")) for i in root.iter("item")]
    assert values == [-4.2, -0.02, 0.02, 4.2]


def test_plan_image_is_bounded_and_draws_the_polygon(tmp_path):
    spec = fixture_spec(0.05)
    write_surface(tmp_path / "t.tif", spec, lambda x, y: plane(x, y) + cone(x, y))
    write_surface(tmp_path / "d.tif", spec, lambda x, y: cone(x, y))
    with open_surface(tmp_path / "t.tif") as top, open_surface(tmp_path / "d.tif") as diff:
        png = render_plan_image(top, diff, circle(CX, CY, 12.0), diff_scale=5.0, max_px=300)
    img = PILImage.open(io.BytesIO(png))
    assert max(img.size) <= 300 and img.mode == "RGB"
    px = img.load()
    from app.volumes.plan_image import POLYGON

    def close(a, b, tol=20):
        return all(abs(int(a[i]) - int(b[i])) <= tol for i in range(3))

    has_polygon_pixel = any(close(px[x, y], POLYGON) for x in range(img.size[0]) for y in range(img.size[1]))
    assert has_polygon_pixel


def test_pdf_names_every_measurement_and_its_numbers(tmp_path):
    items = [_item(), _item(name="Pile 2", fill_m3=1234.5)]
    path = tmp_path / "r.pdf"
    report_pdf.write_report(
        path,
        items,
        title="Site <A>",
        project_name="Kuwait",
        generated_at=datetime(2026, 9, 24, tzinfo=UTC),
        compress=False,
    )
    data = path.read_bytes()
    assert data.startswith(b"%PDF-")
    assert data.count(b"/Type /Page\n") + data.count(b"/Type /Page ") >= 4  # summary, 2 measurements, method
    for text in (b"Pile & 1", b"Pile 2", b"523.6 m\\263", b"1,234.5"):
        assert text in data or text.replace(b"\\263", b"") in data, text
    assert b"page 1/" in data
    # spec §6.2: the areal scale factor is printed with the percentage it implies
    assert b"Areal scale factor" in data and b"0.99962 \\(-0.038 %\\)" in data  # PDF escapes ( )


def test_scale_factor_prints_the_percentage():
    assert report_pdf.scale_factor(1.00031) == "1.00031 (+0.031 %)"
    assert report_pdf.scale_factor(0.99962) == "0.99962 (-0.038 %)"
