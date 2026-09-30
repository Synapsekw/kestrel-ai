"""R5 writers (spec §11): columns, utf-8-sig, severity fills, text safety.

Ruling P4 (index reconciliation item 4): the severity scale comes from
`app.reports.context.read_scale`/`context.Level`, not a local duplicate; `rows._page` reuses R2's
`app.reports.observed.observed_day/host_label/area_m2`.
"""

import csv
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from reports_render_helpers import map_finding

from app.catalogue.handle import DEFAULT_SCALE
from app.reports import context
from app.reports.writers import csv_out, rows, xlsx_out

EXPECTED_COLUMNS = (
    "number, type, type_kind, severity_level, severity_name, status, observed_on, data_item, data_item_type, "
    "anchor_kind, image_file, px_x, px_y, map_x, map_y, map_crs, lon, lat, cloud_x, cloud_y, cloud_z, "
    "uncertainty_m, area_m2, note, created_by, confidence, photos, comments, created_at, updated_at, "
    "report_version"
).split(", ")


def _scale(handle) -> dict[int, context.Level]:
    return {lv.level: lv for lv in context.read_scale(handle)}


@pytest.fixture
def findings(client, handle, project_id, crack):
    made = [
        map_finding(client, handle, project_id, crack["id"], severity=3, note="Crack, 2 m — über the joint"),
        map_finding(client, handle, project_id, crack["id"], severity=None, note="=1+1"),
        map_finding(client, handle, project_id, crack["id"], severity=1),
    ]
    return [f["id"] for f in made]


def test_columns_are_the_spec_list():
    assert rows.COLUMNS == EXPECTED_COLUMNS


def test_rows_follow_the_given_order_and_skip_deleted_ids(handle, findings):
    scale = _scale(handle)
    out = list(rows.export_rows(handle, [findings[2], "gone", findings[0]], scale=scale, version_number=4))
    assert [r["severity_level"] for r in out] == [1, 3]
    first = out[0]
    assert set(first) == set(EXPECTED_COLUMNS)
    assert first["type"] == "crack" and first["type_kind"] == "defect"
    assert first["anchor_kind"] == "map" and first["data_item"] == "April"
    assert first["data_item_type"] == "map"
    assert (first["map_x"], first["map_y"]) == (10.0, 20.0)
    assert first["report_version"] == 4 and first["number"].startswith("F-")
    assert first["severity_name"] == DEFAULT_SCALE[0][1]


def test_ungraded_reads_ungraded(handle, findings):
    out = list(rows.export_rows(handle, [findings[1]], scale=_scale(handle), version_number=1))
    assert out[0]["severity_level"] is None and out[0]["severity_name"] == context.UNGRADED


def test_csv_is_utf8_sig_with_the_exact_header(handle, findings, tmp_path):
    path = tmp_path / "findings.csv"
    n = csv_out.write_csv(path, rows.export_rows(handle, findings, scale=_scale(handle), version_number=1))
    assert n == 3
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
    with path.open(encoding="utf-8-sig", newline="") as f:
        read = list(csv.reader(f))
    assert read[0] == EXPECTED_COLUMNS
    assert read[1][EXPECTED_COLUMNS.index("note")] == "Crack, 2 m — über the joint"


def test_xlsx_sheets_header_filter_and_severity_fill(handle, findings, tmp_path):
    path = tmp_path / "findings.xlsx"
    scale = _scale(handle)
    xlsx_out.write_xlsx(
        path,
        rows.export_rows(handle, findings, scale=scale, version_number=1),
        scale=scale,
        measurements=iter([{c: "" for c in rows.MEASUREMENT_COLUMNS} | {"kind": "volume", "name": "Pile"}]),
        counts=[[["Type", "Count"], ["excavator", 3]]],
        report_info=[("Report", "Weekly inspection"), ("Version", "v001")],
    )
    wb = load_workbook(path)
    assert wb.sheetnames == ["Findings", "Measurements", "Counts", "Report"]
    ws = wb["Findings"]
    assert [c.value for c in ws[1]] == EXPECTED_COLUMNS
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == f"A1:AE{len(findings) + 1}"
    sev = EXPECTED_COLUMNS.index("severity_name") + 1
    level3 = next(lv for lv in DEFAULT_SCALE if lv[0] == 3)
    assert ws.cell(row=2, column=sev).fill.fgColor.rgb == "FF" + level3[2].lstrip("#").upper()
    assert ws.cell(row=3, column=sev).fill.fgColor.rgb == "FF" + xlsx_out.UNGRADED_COLOUR.lstrip("#").upper()
    assert wb["Measurements"]["A1"].value == "kind" and wb["Measurements"]["A2"].value == "volume"
    assert [c.value for c in wb["Counts"][2]] == ["excavator", 3]
    assert wb["Report"]["B1"].value == "Weekly inspection"


def test_xlsx_text_that_starts_with_equals_stays_text(handle, findings, tmp_path):
    path = tmp_path / "findings.xlsx"
    scale = _scale(handle)
    xlsx_out.write_xlsx(
        path,
        rows.export_rows(handle, [findings[1]], scale=scale, version_number=1),
        scale=scale,
        measurements=None,
        counts=[],
        report_info=[],
    )
    cell = load_workbook(path)["Findings"].cell(row=2, column=EXPECTED_COLUMNS.index("note") + 1)
    assert cell.value == "=1+1" and cell.data_type == "s"


def test_read_scale_falls_back_when_the_catalogue_breaks():
    """`rows.py` has no local `severity_scale`/`Level`; the fallback lives in `context.read_scale`
    (Ruling P4). A broken catalogue session still yields the built-in 1-4 scale."""

    class BrokenCatalogue:
        def session(self):
            raise RuntimeError("locked")

    handle = SimpleNamespace(catalogue=BrokenCatalogue())
    levels = {lv.level: lv for lv in context.read_scale(handle)}
    assert levels[1].name == DEFAULT_SCALE[0][1]
