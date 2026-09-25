"""The volume_export job and route (spec 2026-09-23-volumes §10, §11.1 path 17)."""

import csv
import json
import sqlite3

import pytest
from openpyxl import load_workbook
from surfaces import CX, CY, circle, cone, fixture_spec, plane
from volume_rows import add_surface

from app.volumes.paths import diff_path

BASE = "/api/v1/projects"


@pytest.fixture
def project_kind() -> str:
    return "detect"


@pytest.fixture
def measured(client, wait_job, project_id, handle):
    top = add_surface(handle, fixture_spec(0.1), lambda x, y: plane(x, y) + cone(x, y), name="April")
    ids = []
    for name, r in (("Pile north", 12.0), ("Pile north 2", 11.0)):
        body = {
            "name": name,
            "polygon_native": circle(CX, CY, r),
            "top_surface_id": top,
            "base": {"kind": "toe_plane"},
        }
        res = client.post(f"{BASE}/{project_id}/volumes", json=body).json()
        assert wait_job(project_id, res["job"]["id"])["state"] == "succeeded"
        ids.append(res["measurement"]["id"])
    return ids


def _export(client, wait_job, project_id, ids, formats):
    r = client.post(f"{BASE}/{project_id}/volume-exports", json={"measurement_ids": ids, "formats": formats})
    assert r.status_code == 202, r.text
    return wait_job(project_id, r.json()["job"]["id"])


def test_every_format_is_written(client, wait_job, project_id, handle, measured):
    job = _export(client, wait_job, project_id, measured, ["pdf", "gpkg", "csv", "xlsx"])
    assert job["state"] == "succeeded", job
    folder = handle.folder / job["result"]["folder"]
    files = job["result"]["files"]
    assert files == [
        "volumes-report.pdf",
        "volumes.gpkg",
        "pile-north-cutfill.tif",
        "pile-north-cutfill.qml",
        "pile-north-2-cutfill.tif",
        "pile-north-2-cutfill.qml",
        "volumes.csv",
        "volumes.xlsx",
        "summary.json",
    ]
    assert (folder / "volumes-report.pdf").read_bytes().startswith(b"%PDF-")
    con = sqlite3.connect(folder / "volumes.gpkg")
    srs = {r[0] for r in con.execute("SELECT srs_id FROM gpkg_spatial_ref_sys")}
    assert 32639 in srs
    assert con.execute('SELECT count(*) FROM "measurements"').fetchone()[0] == 2
    con.close()
    assert (folder / "pile-north-cutfill.tif").read_bytes() == diff_path(handle, measured[0]).read_bytes()
    rows = list(csv.DictReader((folder / "volumes.csv").open(encoding="utf-8")))
    assert [r["name"] for r in rows] == ["Pile north", "Pile north 2"]
    assert load_workbook(folder / "volumes.xlsx").sheetnames == ["Volumes", "Method", "Warnings"]
    summary = json.loads((folder / "summary.json").read_text("utf-8"))
    assert [m["name"] for m in summary["measurements"]] == ["Pile north", "Pile north 2"]
    stored = client.get(f"{BASE}/{project_id}/volumes/{measured[0]}").json()["results"]["fill_m3"]
    assert float(rows[0]["fill_m3"]) == pytest.approx(stored)


def test_stale_or_unknown_measurements_are_refused(client, project_id, measured):
    url = f"{BASE}/{project_id}/volume-exports"
    assert client.post(url, json={"measurement_ids": ["nope"], "formats": ["csv"]}).status_code == 404
    client.patch(f"{BASE}/{project_id}/volumes/{measured[0]}", json={"base": {"kind": "flat", "z": 50.0}})
    r = client.post(url, json={"measurement_ids": measured, "formats": ["csv"]})
    assert r.status_code == 409, r.text
    error = r.json()["error"]
    assert error["code"] == "not_ready"
    assert "recalculate" in error["message"]


def test_the_title_defaults_to_the_project_name(client, wait_job, project_id, measured):
    job = _export(client, wait_job, project_id, measured[:1], ["csv"])
    assert job["params"]["title"] == "T"
