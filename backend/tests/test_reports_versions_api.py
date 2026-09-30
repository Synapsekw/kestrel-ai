"""R5 endpoints (spec §14) with the real R1-R4 pipeline for the end-to-end case."""

import csv
import threading
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook
from reports_render_helpers import API, map_finding, new_report, start_render

from app.reports import render_job, versions
from app.reports.models import ReportVersion
from app.reports.theme import THEME_VERSION


@pytest.fixture
def live_render(app, monkeypatch):
    monkeypatch.setattr(render_job, "run_pipeline", render_job._run_pipeline)


@pytest.fixture
def report(client, project_id):
    return new_report(client, project_id)


def _base(project_id, report_id):
    return f"{API}/projects/{project_id}/reports/{report_id}"


def test_a_real_render_produces_v1_with_pdf_csv_xlsx(
    live_render, client, handle, project_id, report, crack, wait_job
):
    for sev in (1, 3, None):
        map_finding(client, handle, project_id, crack["id"], severity=sev)
    job = start_render(client, project_id, report["id"])
    done = wait_job(project_id, job["id"])
    assert done["state"] == "succeeded", done
    page = client.get(f"{_base(project_id, report['id'])}/versions").json()
    v = page["items"][0]
    assert (v["number"], v["state"]) == (1, "ready")
    kinds = sorted(f["kind"] for f in v["files"])
    assert kinds == ["csv", "pdf", "xlsx"]
    folder = handle.folder / v["folder"]
    pdf = next(folder.glob("*-v001*.pdf"))
    assert pdf.read_bytes().startswith(b"%PDF")
    with (folder / "findings.csv").open(encoding="utf-8-sig", newline="") as f:
        assert len(list(csv.reader(f))) == 1 + 3
    assert load_workbook(folder / "findings.xlsx").sheetnames[0] == "Findings"
    assert (folder / "document.json").is_file()


def test_second_render_of_the_same_report_is_409(client, project_id, report, monkeypatch):
    gate = threading.Event()

    def slow(ctx):
        gate.wait(10)
        raise render_job.JobFailure("stopped by the test")

    monkeypatch.setattr(render_job, "run_pipeline", slow)
    try:
        start_render(client, project_id, report["id"])
        r = client.post(f"{_base(project_id, report['id'])}/renders", json={"formats": ["pdf"]})
        assert r.status_code == 409 and r.json()["error"]["code"] == "render_running"
    finally:
        gate.set()


def test_render_of_an_unknown_report_is_404(client, project_id):
    r = client.post(f"{_base(project_id, 'nope')}/renders", json={"formats": ["pdf"]})
    assert r.status_code == 404


def test_a_failed_render_is_listed_without_a_number(client, project_id, report, wait_job):
    job = start_render(client, project_id, report["id"])  # the conftest seam fails it
    assert wait_job(project_id, job["id"])["state"] == "failed"
    v = client.get(f"{_base(project_id, report['id'])}/versions").json()["items"][0]
    assert (v["state"], v["number"]) == ("failed", None)
    assert "disabled in tests" in v["stats"]["error"]


def _ready_version(handle, report_id, n=1):
    vid = versions.create_rendering_row(handle, report_id, label=None)
    partial = versions.reports_root(handle, report_id) / f".partial-2026-09-30_10000{n}"
    partial.mkdir(parents=True)
    (partial / "document.json").write_text(
        render_job.ReportDocument.model_validate(
            {
                "report_id": report_id,
                "version": n,
                "generated_at": "2026-09-30T10:00:00Z",
                "theme_version": str(THEME_VERSION),
                "sections": [
                    {
                        "key": "summary",
                        "title": "Summary",
                        "blocks": [{"kind": "para", "text": f"p{i}", "style": "body"} for i in range(60)],
                    }
                ],
            }
        ).model_dump_json(),
        "utf-8",
    )
    versions.promote(
        handle,
        report_id=report_id,
        version_id=vid,
        number=n,
        partial=partial,
        files=[],
        stats={},
        baseline_version_id=None,
        states=[],
    )
    with handle.session() as s:  # distinct times: Windows' clock can give two rows the same instant
        s.get(ReportVersion, vid).created_at = datetime(2026, 9, 30, 10, n, tzinfo=UTC)
    return vid


def test_issue_unissue_and_delete_rules(client, handle, project_id, report):
    _ready_version(handle, report["id"])
    base = f"{_base(project_id, report['id'])}/versions/1"
    r = client.patch(base, json={"issued": True})
    assert r.status_code == 200 and r.json()["issued_at"] is not None
    r = client.delete(base)
    assert r.status_code == 409 and r.json()["error"]["code"] == "issued_version"
    assert client.patch(base, json={"issued": False}).json()["issued_at"] is None
    assert client.delete(base).status_code == 204
    assert client.get(base).status_code == 404
    assert not (versions.reports_root(handle, report["id"]) / "v001").exists()


def test_detail_has_files_and_unknown_number_is_404(client, handle, project_id, report):
    _ready_version(handle, report["id"])
    assert client.get(f"{_base(project_id, report['id'])}/versions/1").json()["number"] == 1
    assert client.get(f"{_base(project_id, report['id'])}/versions/9").status_code == 404


def test_document_pages_by_fifty(client, handle, project_id, report):
    _ready_version(handle, report["id"])
    url = f"{_base(project_id, report['id'])}/versions/1/document"
    first = client.get(url).json()
    assert sum(len(s["blocks"]) for s in first["sections"]) == 50 and first["next_cursor"]
    second = client.get(url, params={"cursor": first["next_cursor"]}).json()
    assert [b["text"] for b in second["sections"][0]["blocks"]] == [f"p{i}" for i in range(50, 60)]
    assert second["next_cursor"] is None
    assert client.get(url, params={"limit": 51}).status_code == 422


def test_versions_list_pages_newest_first(client, handle, project_id, report):
    for n in (1, 2, 3):
        _ready_version(handle, report["id"], n)
    base = f"{_base(project_id, report['id'])}/versions"
    first = client.get(base, params={"limit": 2}).json()
    assert [v["number"] for v in first["items"]] == [3, 2] and first["next_cursor"]
    rest = client.get(base, params={"cursor": first["next_cursor"]}).json()
    assert [v["number"] for v in rest["items"]] == [1]
