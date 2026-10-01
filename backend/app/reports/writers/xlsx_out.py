"""findings.xlsx (spec §11.2): openpyxl write-only; sheets Findings (frozen header, autofilter,
severity cell filled with the scale colour), Measurements, Counts, Report. openpyxl loads inside
`write_xlsx` only.

The ungraded fill reuses R2's grey (`app.reports.context.GREY`) so the PDF and the XLSX agree
(Ruling P4)."""

from __future__ import annotations

from collections.abc import Callable, Iterable
from pathlib import Path

from app.reports.context import GREY
from app.reports.writers.rows import COLUMNS, MEASUREMENT_COLUMNS, Level

UNGRADED_COLOUR = GREY


def _argb(colour: str) -> str:
    return "FF" + colour.lstrip("#").upper()


def write_xlsx(
    path: Path,
    rows: Iterable[dict],
    *,
    scale: dict[int, Level],
    measurements: Iterable[dict] | None,
    counts: list[list[list]],
    report_info: list[tuple[str, str]],
    on_row: Callable[[int], None] | None = None,
) -> int:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    def cell(ws, value, *, font=None, fill=None):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)  # one control character must not fail the render
        c = WriteOnlyCell(ws, value=value)
        if isinstance(value, str) and value.startswith("="):
            c.data_type = "s"  # a note like "=1+1" is text, never a formula (Review Focus 4)
        if font is not None:
            c.font = font
        if fill is not None:
            c.fill = fill
        return c

    def header(ws, names):
        ws.append([cell(ws, n, font=bold) for n in names])

    findings = wb.create_sheet("Findings")
    findings.freeze_panes = "A2"
    header(findings, COLUMNS)
    sev = COLUMNS.index("severity_name")
    n = 0
    for row in rows:
        values = [cell(findings, row[c]) for c in COLUMNS]
        level = scale.get(row["severity_level"]) if row["severity_level"] is not None else None
        colour = level.colour if level is not None else UNGRADED_COLOUR
        values[sev] = cell(findings, row["severity_name"], fill=PatternFill("solid", fgColor=_argb(colour)))
        findings.append(values)
        n += 1
        if on_row is not None:
            on_row(n)
    findings.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{n + 1}"

    if measurements is not None:
        ws = wb.create_sheet("Measurements")
        ws.freeze_panes = "A2"
        header(ws, MEASUREMENT_COLUMNS)
        for m in measurements:
            ws.append([cell(ws, m.get(c)) for c in MEASUREMENT_COLUMNS])

    if counts:
        ws = wb.create_sheet("Counts")
        for t, table in enumerate(counts):
            if t:
                ws.append([])
            head, *body = table
            header(ws, head)
            for r in body:
                ws.append([cell(ws, v) for v in r])

    ws = wb.create_sheet("Report")
    for label, value in report_info:
        ws.append([cell(ws, label, font=bold), cell(ws, value)])

    wb.save(path)
    return n
