"""`kestrel-backend.exe reports-selftest`: proves the frozen bundle carries what `report_render` loads
lazily (spec 2026-09-26-reports §17): the TTF fonts (sha256-checked, registered, not the Helvetica
fallback), reportlab's linear gradient, JPEG passthrough, a reportlab.graphics chart, and openpyxl in
write-only mode with a fill. smoke_frozen.ps1 runs it after volumes-selftest and expects "reports ok"."""

from __future__ import annotations

import tempfile
from datetime import UTC, datetime
from pathlib import Path


def run() -> dict:
    from openpyxl import Workbook, load_workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import PatternFill
    from PIL import Image as PILImage
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph

    from app.reports.pdf import canvas, charts, fonts, primitives, styles

    fs = fonts.register_fonts()
    st = styles.build_styles(fs)
    with tempfile.TemporaryDirectory(prefix="kestrel-reports-") as tmp:
        work = Path(tmp)
        jpg = work / "snap.jpg"
        PILImage.new("RGB", (320, 240), (90, 120, 160)).save(jpg, "JPEG", quality=85)
        pdf = work / "selftest.pdf"
        meta = canvas.PageMeta(
            title="Kestrel AI reports selftest",
            version_label="v1",
            project="Selftest",
            generated_at=datetime(2026, 9, 30, tzinfo=UTC),
            cover_pages=frozenset({1}),
        )
        doc = BaseDocTemplate(str(pdf), pagesize=A4, invariant=1, initialFontName=fs.sans)
        below = A4[1] - canvas.band_height(A4[1])
        frame = Frame(18 * mm, 18 * mm, A4[0] - 36 * mm, below - 24 * mm)
        on_page = lambda c, _d: canvas.draw_cover_band(c, None)  # noqa: E731
        doc.addPageTemplates([PageTemplate("cover", [frame], onPage=on_page)])
        doc.build(
            [
                Paragraph("Reports selftest m³", st.h1),
                primitives.figure_flowable(jpg, 80 * mm, 60 * mm, "JPEG passthrough", st),
                charts.chart_drawing(
                    "bar",
                    [charts.ChartSeries("Findings", [3, 5, 2])],
                    ["A", "B", "C"],
                    "count",
                    width=120 * mm,
                    height=60 * mm,
                    styles=st,
                ),
            ],
            canvasmaker=canvas.canvas_class(meta, st),
        )
        data = pdf.read_bytes()
        xlsx = work / "selftest.xlsx"
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Findings")
        cell = WriteOnlyCell(ws, value="Major")
        cell.fill = PatternFill("solid", fgColor="FFE08A")
        ws.append(["Severity"])
        ws.append([cell])
        wb.save(xlsx)
        back = load_workbook(xlsx)["Findings"]["A2"]
        return {
            "fonts": fonts.verify_fonts() if fs.embedded else 0,
            "gradient": data.count(b"/ShadingType 2"),
            "jpeg": data.count(b"/DCTDecode"),
            "chart": 1 if data.count(b"/Subtype /Image") == 1 else 0,  # the JPEG only: the chart is vector
            "xlsx": back.value,
            "fill": str(back.fill.fgColor.rgb),
            "pdf_bytes": len(data),
        }


def main(argv: list[str] | None = None) -> int:
    """`argv` (the arguments after `reports-selftest`) is accepted and ignored."""
    f = run()
    ok = (
        f["fonts"] == 3
        and f["gradient"] >= 1
        and f["jpeg"] == 1
        and f["chart"] == 1
        and f["xlsx"] == "Major"
        and f["fill"].endswith("FFE08A")
    )
    print(
        f"reports {'ok' if ok else 'FAILED'} fonts {f['fonts']} gradient {f['gradient']} jpeg {f['jpeg']} "
        f"chart {f['chart']} xlsx {f['xlsx']} pdf {f['pdf_bytes']}"
    )
    return 0 if ok else 1
