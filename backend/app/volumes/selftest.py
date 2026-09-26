"""`kestrel-backend.exe volumes-selftest`: proves the frozen bundle carries what `volume_export`
and `volume_calc` load lazily - reportlab (platypus, the built-in Helvetica), openpyxl in write-only
mode, and scipy.spatial's qhull behind Delaunay and LinearNDInterpolator (plan Task 13)."""

from __future__ import annotations

import io
import tempfile
from pathlib import Path


def run() -> dict:
    import numpy as np
    from openpyxl import Workbook, load_workbook
    from openpyxl.cell import WriteOnlyCell
    from PIL import Image as PILImage
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Table
    from scipy.interpolate import LinearNDInterpolator
    from scipy.spatial import Delaunay

    with tempfile.TemporaryDirectory(prefix="kestrel-volumes-") as tmp:
        pdf = Path(tmp) / "selftest.pdf"
        png = io.BytesIO()
        PILImage.new("RGB", (64, 32), (200, 180, 140)).save(png, "PNG")
        png.seek(0)
        styles = getSampleStyleSheet()
        SimpleDocTemplate(str(pdf), pagesize=A4).build(
            [
                Paragraph("Kestrel AI volumes selftest", styles["Title"]),
                Table([["fill", "523.6 m³"]]),
                Image(png, 64, 32),
            ]
        )
        xlsx = Path(tmp) / "selftest.xlsx"
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Volumes")
        cell = WriteOnlyCell(ws, value=523.6)
        cell.number_format = "#,##0.0"
        ws.append(["fill_m3"])
        ws.append([cell])
        wb.save(xlsx)
        value = load_workbook(xlsx)["Volumes"]["A2"].value
        pts = np.array([[0.0, 0.0], [1.0, 0.0], [0.0, 1.0], [1.0, 1.0]])
        tri = Delaunay(pts)
        z = float(LinearNDInterpolator(tri, pts[:, 0] + 2 * pts[:, 1])(0.25, 0.5))
        return {"pdf_bytes": pdf.stat().st_size, "xlsx_value": value, "simplices": len(tri.simplices), "z": z}


def main(argv: list[str] | None = None) -> int:
    facts = run()
    ok = facts["pdf_bytes"] > 500 and facts["xlsx_value"] == 523.6 and abs(facts["z"] - 1.25) < 1e-9
    verdict = "ok" if ok else "FAILED"
    print(
        f"volumes {verdict} pdf {facts['pdf_bytes']} xlsx {facts['xlsx_value']} delaunay {facts['simplices']}"
    )
    return 0 if ok else 1
